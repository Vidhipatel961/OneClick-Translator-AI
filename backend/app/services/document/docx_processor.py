import os
import uuid
from typing import List, Callable, Optional
from docx import Document
from app.services.document.base import DocumentProcessor
from app.core.exceptions import BaseLingoraException

class DOCXProcessor(DocumentProcessor):
    def validate(self, file_path: str):
        if not os.path.exists(file_path):
            raise BaseLingoraException("File not found.")
        size_mb = os.path.getsize(file_path) / (1024 * 1024)
        if size_mb > 50:
            raise BaseLingoraException("DOCX size exceeds 50MB limit.")
            
    def _iter_paragraphs_and_cells(self, doc):
        """Yield all paragraphs from body and tables."""
        for p in doc.paragraphs:
            yield p
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        yield p

    def extract_text(self, file_path: str, progress_callback: Optional[Callable[[str, int, int], None]] = None, lang: str = "en") -> List[str]:
        self.validate(file_path)
        doc = Document(file_path)
        chunks = []
        for p in self._iter_paragraphs_and_cells(doc):
            if p.text.strip():
                chunks.append(p.text)
        return chunks
        
    def generate_translated_file(self, original_file_path: str, translated_texts: List[str], output_dir: str) -> str:
        doc = Document(original_file_path)
        
        trans_idx = 0
        # Iterate in the EXACT same order to replace text while preserving style on the first run of the paragraph
        for p in self._iter_paragraphs_and_cells(doc):
            if p.text.strip():
                if trans_idx < len(translated_texts):
                    # We clear runs and set text on the first run to roughly preserve paragraph style
                    if len(p.runs) > 0:
                        style = p.runs[0].style
                        p.text = translated_texts[trans_idx]
                        if len(p.runs) > 0:
                            p.runs[0].style = style
                    else:
                        p.text = translated_texts[trans_idx]
                    trans_idx += 1
                    
        os.makedirs(output_dir, exist_ok=True)
        output_path = os.path.join(output_dir, f"translated_{uuid.uuid4().hex[:8]}.docx")
        doc.save(output_path)
        return output_path
